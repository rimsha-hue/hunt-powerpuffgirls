"""Generate fixtures/demo_resume.docx, the fictional resume behind the demo account and "Try with sample data".

Run: python scripts/make_demo_resume.py
"""

from pathlib import Path

import docx
from docx.shared import Pt

OUT = Path(__file__).resolve().parents[1] / "fixtures" / "demo_resume.docx"

SECTIONS: list[tuple[str, list[str]]] = [
    (
        "SUMMARY",
        [
            "Backend and AI engineer who builds reliable Python services and LLM features that ship.",
        ],
    ),
    (
        "EXPERIENCE",
        [
            "Machine Learning Engineer, Nimbus Labs (Pvt.) Ltd., Lahore, Jul 2023 - Present",
            "• Built a retrieval-augmented generation (RAG) assistant with LangChain and Qdrant used by 1,200 support agents",
            "• Reduced LLM inference cost by 40% by adding Redis response caching and prompt compression",
            "• Shipped FastAPI microservices on Docker and AWS ECS with GitHub Actions CI/CD",
            "Software Engineer Intern, Acme Analytics, Islamabad, Jun 2022 - Aug 2022",
            "• Automated ETL pipelines in Apache Airflow processing 2M records daily",
            "• Cut dashboard query time by 3x by adding PostgreSQL indexes",
        ],
    ),
    (
        "PROJECTS",
        [
            "ResumeRadar: NLP pipeline that ranks resumes against job descriptions using spaCy and scikit-learn",
            "UrduSentiment: fine-tuned a BERT model with PyTorch for Urdu sentiment analysis, reaching 87% accuracy",
        ],
    ),
    (
        "EDUCATION",
        [
            "BS Computer Science, National University of Sciences and Technology (NUST), 2019 - 2023 | CGPA 3.6/4.0",
        ],
    ),
    (
        "CERTIFICATIONS",
        [
            "AWS Certified Cloud Practitioner, Amazon Web Services, 2024",
            "Deep Learning Specialization, DeepLearning.AI, 2022",
        ],
    ),
    (
        "SKILLS",
        [
            "Languages: Python, SQL, TypeScript, JavaScript",
            "AI/ML: PyTorch, scikit-learn, LangChain, Hugging Face, RAG, Prompt Engineering",
            "Backend: FastAPI, PostgreSQL, Redis, Docker, AWS, Git",
            "Frontend: React, Tailwind CSS",
        ],
    ),
]


def main() -> None:
    document = docx.Document()
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    document.add_heading("AYESHA KHAN", level=0)
    document.add_paragraph(
        "ayesha.khan@example.com | +92 300 0000000 | linkedin.com/in/ayesha-khan-demo | github.com/ayesha-khan-demo"
    )
    document.add_paragraph("Lahore, Pakistan")
    for heading, lines in SECTIONS:
        document.add_heading(heading, level=1)
        for line in lines:
            document.add_paragraph(line)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(OUT))
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
