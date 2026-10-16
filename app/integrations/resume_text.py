"""Text extraction from PDF and DOCX resumes. CPU-bound, so it runs in a worker thread."""

import io

import anyio
import structlog

from app.domain.uploads import FileKind, UploadRejectedError

logger = structlog.get_logger()

MAX_PDF_PAGES = 10


def _pdf_text(data: bytes) -> str:
    import pdfplumber

    with pdfplumber.open(io.BytesIO(data)) as pdf:
        pages = pdf.pages[:MAX_PDF_PAGES]
        return "\n".join((page.extract_text() or "") for page in pages)


def _docx_text(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells if cell.text.strip()))
    return "\n".join(parts)


def extract_text(data: bytes, kind: FileKind) -> str:
    try:
        return _pdf_text(data) if kind == "pdf" else _docx_text(data)
    except Exception as exc:
        logger.warning("resume_text_extraction_failed", kind=kind, error_type=type(exc).__name__)
        raise UploadRejectedError(
            "We couldn't read that file. Try exporting it again as PDF or DOCX."
        ) from exc


async def extract_text_async(data: bytes, kind: FileKind) -> str:
    return await anyio.to_thread.run_sync(extract_text, data, kind)
